"""Inspect the frozen Germany stress magnitude package without rewriting it."""
import io
import zipfile
from collections import Counter
import openpyxl
from pypdf import PdfReader
from workspace import ROOT, target, write_csv, write_json

def main():
    with zipfile.ZipFile(ROOT/'sources/new/N08/stress_magnitudes.zip') as z:
        wb=openpyxl.load_workbook(io.BytesIO(z.read('stressmagdata_germany_2020.xlsx')),read_only=True,data_only=True)
        inventories=[];records=[]
        for ws in wb.worksheets:
            rows=list(ws.iter_rows(values_only=True))
            inventories.append(dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column,first_rows=[list(r)[:12] for r in rows[:7]]))
            header_index=next((i for i,r in enumerate(rows) if str(r[0]).strip().lower()=='id' and 'LAT' in r),None)
            if header_index is not None:
                header=[str(x) if x is not None else f'unnamed_{i}' for i,x in enumerate(rows[header_index])]
                records.extend({k:v for k,v in zip(header,r)} for r in rows[header_index+1:] if r[0] is not None)
        wb.close()
        pdf=z.read('List_of_parameters.pdf')
        target('output/extracted/N08_List_of_parameters.pdf').write_bytes(pdf)
        text='\n'.join(f'=== PDF page {i+1} ===\n'+(p.extract_text() or '') for i,p in enumerate(PdfReader(io.BytesIO(pdf)).pages))
        target('output/extracted/N08_parameters.txt').write_text(text,encoding='utf-8')
    keys=list(records[0]) if records else []
    result=dict(workbooks=inventories,rows=len(records),fields=keys,
                missing={k:sum(r.get(k) in (None,'') for r in records) for k in keys},
                categories={k:dict(Counter(str(r.get(k)) for r in records)) for k in ['REG','TYPE','QUALITY','DEPTH_REF']},
                scope='Local contemporary stress magnitudes; not global force-source counts, lifetimes, or a rift-only population.')
    write_csv('output/tables/Germany_stress_magnitudes.csv',records)
    write_json('output/checks/Germany_magnitude_profile.json',result)
    print(result['rows'],keys)
    print(result['categories'])

if __name__=='__main__': main()
