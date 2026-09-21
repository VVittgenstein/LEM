"""Read-only source profiling. All generated evidence stays under rifting/output."""
from __future__ import annotations
import csv
import io
import json
import math
import statistics
import zipfile
from collections import Counter
from pathlib import Path
from workspace import ROOT, digest, read_csv, target, utc_now, write_csv, write_json

def counts(values):
    return dict(sorted(Counter(str(v) if v is not None else '<missing>' for v in values).items()))

def numeric(value):
    try:
        out = float(value)
        return out if math.isfinite(out) else None
    except (TypeError, ValueError):
        return None

def summary(values):
    vals = sorted(x for v in values if (x := numeric(v)) is not None)
    return dict(n=len(vals), minimum=min(vals) if vals else None,
                median=statistics.median(vals) if vals else None,
                maximum=max(vals) if vals else None)

def basic(rows, key=None):
    fields = list(rows[0]) if rows else []
    out = dict(rows=len(rows), fields=fields,
               missing={k:sum(r.get(k) in ('',None) for r in rows) for k in fields},
               exact_duplicate_rows=len(rows)-len({json.dumps(r,sort_keys=True,ensure_ascii=False) for r in rows}))
    if key:
        c=Counter(str(r.get(key)) for r in rows)
        out.update(key=key, unique_keys=len(c), duplicate_keys={k:n for k,n in c.items() if n>1})
    return out

def read_geo(path):
    return json.loads(path.read_text(encoding='utf-8'))['features']

def csv_profile():
    existing=ROOT/'sources/existing/tables'
    catalogue=read_csv(existing/'T01_rift_catalogue.csv')
    t=basic(catalogue,'catalogue_record_id')
    for f in ['width_qualifier','entry_hierarchy','onset_assignment','end_age_blank','independent_event_status']:
        t[f]=counts(r[f] for r in catalogue)
    t['length_km']=summary(r['length_numeric_original'] for r in catalogue)
    t['width_km']=summary(r['width_numeric_original'] for r in catalogue)
    t['use']='Regional structures and catalogue dimensions; not independent fault episodes or a duration sample.'
    wsm_path=ROOT/'sources/new/N01/WSM_Database_2025.csv'
    wsm=read_csv(wsm_path)
    w=basic(wsm,'ID')
    w['quality']=counts(r['QUALITY'] for r in wsm)
    w['regime']=counts(r['REGIME'] for r in wsm)
    w['indicator_type']=counts(r['TYPE'] for r in wsm)
    ac=[r for r in wsm if r['QUALITY'] in ('A','B','C')]
    nf=[r for r in ac if r['REGIME']=='NF']
    w['abc_count']=len(ac)
    w['abc_normal_fault_regime_count']=len(nf)
    w['abc_nf_indicator_types']=counts(r['TYPE'] for r in nf)
    w['azimuth_valid_0_180']=sum(numeric(r['AZI']) is not None and 0<=float(r['AZI'])<=180 for r in wsm)
    w['abc_nf_depth_km']=summary(r['DEPTH'] for r in nf)
    w['latitude_out_of_range']=sum(numeric(r['LAT']) is None or abs(float(r['LAT']))>90 for r in wsm)
    w['longitude_out_of_range']=sum(numeric(r['LON']) is None or abs(float(r['LON']))>180 for r in wsm)
    w['use']='Contemporary stress indicators; AZI is axial SHmax orientation. EQ_MAG is earthquake magnitude, not stress magnitude. No force-source history.'
    write_csv('output/tables/WSM_ABC_NF.csv', [{k:r[k] for k in ['ID','LAT','LON','AZI','TYPE','DEPTH','QUALITY','REGIME','REF1','PLATE']} for r in nf])
    rate_rows=read_csv(existing/'R05_growth_strata_fault_rates.csv')
    rate=basic(rate_rows)
    rate['fault_ids']=counts(r['fault_id'] for r in rate_rows)
    rate['quantity_status']=counts(r['quantity_status'] for r in rate_rows)
    rate['author_rate_m_per_Myr']=summary(r['author_growth_rate_m_per_Myr'] for r in rate_rows)
    errors=[]
    for row in rate_rows:
        calculated=(float(row['hangingwall_thickness_km'])-float(row['footwall_thickness_km']))*1000/(float(row['old_age_Ma'])-float(row['young_age_Ma']))
        errors.append(abs(calculated-float(row['author_growth_rate_m_per_Myr'])))
    rate['independent_recalculation_max_abs_error_m_per_Myr']=max(errors)
    profiles={}
    for name in ['R01_Basement_Lake_Turkana.csv','R01_Gombe_Lake_Turkana.csv','B07_backstripping_profile_rates.csv','activity_records.csv']:
        rows=read_csv(existing/name)
        profiles[name]=basic(rows)
    return dict(T01=t, WSM2025=w, R05=rate, existing_other=profiles)

def geo_profile():
    features=read_geo(ROOT/'sources/existing/sources/S01/gem_active_faults_harmonized.geojson')
    properties=[f['properties'] for f in features]
    normals=[(i,f) for i,f in enumerate(features) if 'normal' in str(f['properties'].get('slip_type','')).lower()]
    normal_rows=[]
    selected=['catalog_id','catalog_name','name','slip_type','average_dip','average_rake','upper_seis_depth','lower_seis_depth','net_slip_rate','vert_sep_rate','shortening_rate','strike_slip_rate']
    for i,f in normals:
        normal_rows.append(dict(source_feature_index=i, **{k:f['properties'].get(k) for k in selected}))
    write_csv('output/tables/GEM_normal_label_attributes.csv',normal_rows)
    g=basic(normal_rows,'catalog_id')
    g['all_features']=len(features)
    g['all_slip_types']=counts(p.get('slip_type') for p in properties)
    g['normal_selection']='case-insensitive normal substring in slip_type; excludes reverse and unclassified; mixed kinematics retained'
    g['catalogues']=counts(row['catalog_name'] for row in normal_rows)
    g['nonempty_dip_tuples']=counts(row['average_dip'] for row in normal_rows if row['average_dip'])
    g['use']='Attributes may be measured, inferred or defaults. Map features are not independent tectonic events.'
    m={}
    for category in ['faults','sections','multifaults']:
        fs=read_geo(ROOT/f'sources/new/N02/MSSM_{category}.geojson')
        rows=[f['properties'] for f in fs]
        assert all('MSSM_id' in row for row in rows)
        p=basic(rows,'MSSM_id')
        p['sample']=rows[0]
        for f in ['slip_rate','length','dip_int','class','basin']:
            p[f]=summary(r.get(f) for r in rows) if f in ['slip_rate','length','dip_int'] else counts(r.get(f) for r in rows)
        p['dip_triplets']=counts(tuple(numeric(r.get(k)) for k in ['dip_lower','dip_int','dip_upper']) for r in rows)
        p['strike_degrees']=summary(r.get('strike') for r in rows)
        m[category]=p
    return dict(GEM=g, MSSM=m)

def workbook_profile():
    import openpyxl
    archive=ROOT/'sources/new/N03/fault_statistics.zip'
    results=[]
    with zipfile.ZipFile(archive) as z:
        write_json('output/checks/Pan_archive_inventory.json',[dict(name=i.filename,bytes=i.file_size) for i in z.infolist()])
        for name in z.namelist():
            if not name.endswith('.xlsx') or '/._' in name or name.startswith('__MACOSX/'):
                continue
            wb=openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True,data_only=True)
            sheets=[]
            for ws in wb.worksheets:
                first=[];nonnull=0;scalar_rows=[];header=None
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        nonnull+=1
                        if len(first)<3: first.append([str(v)[:200] if isinstance(v,str) else v for v in list(row)[:14]])
                        if header is None:
                            header=[str(v) if v is not None else f'unnamed_{i}' for i,v in enumerate(row)]
                        else:
                            scalar_rows.append({header[i]:v for i,v in enumerate(row) if i<len(header) and (not isinstance(v,str) or len(v)<100)})
                scalar_fields={k:summary(r.get(k) for r in scalar_rows) for k in header if k in ['length','dip','throw','timestep','age','max d','sum strain','strike']}
                sheets.append(dict(sheet=ws.title,max_row=ws.max_row,max_column=ws.max_column,nonempty_rows=nonnull,first_nonempty_rows=first,scalar_fields=scalar_fields))
                if '/Observational_data/' in name or '/Geometric_relationships/' in name:
                    keys=list(dict.fromkeys(k for r in scalar_rows for k in r))
                    write_csv('output/tables/Pan_'+Path(name).stem+'.csv',[{k:r.get(k) for k in keys} for r in scalar_rows])
            results.append(dict(member=name,kind='natural_observation' if '/Observational_data/' in name else 'numerical_model',sheets=sheets))
            wb.close()
            print('Inspected',name,flush=True)
            write_json('output/checks/Pan_workbooks.json',results)
    return results

def pdf_evidence():
    from pypdf import PdfReader
    output=[]
    for filename in ['sources/new/N01/WSM_TR_25_01.pdf','sources/new/N03/supplement.pdf']:
        reader=PdfReader(ROOT/filename)
        texts=[]
        for i,page in enumerate(reader.pages):
            texts.append(f'\n=== PDF page {i+1} ===\n'+(page.extract_text() or ''))
        dest=target('output/extracted/'+Path(filename).stem+'.txt')
        dest.write_text('\n'.join(texts),encoding='utf-8')
        output.append(dict(source=filename,pages=len(reader.pages),text=dest.relative_to(ROOT).as_posix(),sha256=digest(ROOT/filename)))
    return output

def main():
    import sys
    result=dict(created_utc=utc_now(),csv=csv_profile(),geo=geo_profile())
    if '--tables-only' in sys.argv:
        prior=json.loads((ROOT/'output/checks/source_profiles.json').read_text(encoding='utf-8'))
        result.update({k:prior[k] for k in ['pdf','workbooks']})
        write_json('output/checks/source_profiles.json',result)
        print('Table profiles updated; unchanged PDF and workbook checks retained.')
        return
    write_json('output/checks/source_profiles.json',result)
    result['pdf']=pdf_evidence()
    result['workbooks']=workbook_profile()
    write_json('output/checks/source_profiles.json',result)
    print('Source profiling complete.',flush=True)

if __name__=='__main__': main()
